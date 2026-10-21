import csv, json
from pathlib import Path
import pandas as pd
import pdfplumber
from PIL import Image
import pytesseract
from sqlalchemy.orm import Session
from app.models.models import UploadedFile, Record, AuditLog
from app.services.cleaning import suggest_mapping, clean_record
from app.services.dedup import should_compare, score_records

def block_keys(data: dict):
    keys=[]
    for field in ("phone", "email", "gst_number"):
        value=str(data.get(field) or "").strip().lower()
        if value: keys.append((field,value))
    name=str(data.get("name") or "").strip().lower()
    if name: keys.append(("initial",name[0]))
    return keys

def read_rows(path: Path):
    ext = path.suffix.lower()
    if ext == ".csv":
        try: return pd.read_csv(path, dtype=str, chunksize=10000)
        except UnicodeDecodeError: return pd.read_csv(path, dtype=str, encoding="latin-1", chunksize=10000)
    if ext == ".xlsx":
        def xlsx_chunks():
            from openpyxl import load_workbook
            workbook=load_workbook(path,read_only=True,data_only=True)
            try:
                iterator=workbook.active.iter_rows(values_only=True)
                headers=next(iterator,None)
                if headers is None: return
                headers=[str(value).strip() if value is not None and str(value).strip() else f"column_{i+1}" for i,value in enumerate(headers)]
                while batch:=list(__import__("itertools").islice(iterator,10000)):
                    yield pd.DataFrame(batch,columns=headers).astype("string").where(lambda x:x.notna(),None)
            finally: workbook.close()
        return xlsx_chunks()
    if ext == ".xls":
        def xls_chunks():
            import xlrd
            workbook=xlrd.open_workbook(path,on_demand=True)
            try:
                sheet=workbook.sheet_by_index(0)
                if sheet.nrows==0:return
                headers=[str(value).strip() if value is not None and str(value).strip() else f"column_{i+1}" for i,value in enumerate(sheet.row_values(0))]
                for start in range(1,sheet.nrows,10000):
                    yield pd.DataFrame([sheet.row_values(i) for i in range(start,min(start+10000,sheet.nrows))],columns=headers).astype("string").where(lambda x:x.notna(),None)
            finally: workbook.release_resources()
        return xls_chunks()
    text = ""
    if ext == ".pdf":
        with pdfplumber.open(path) as pdf:
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
            if not text.strip():
                text = "\n".join(pytesseract.image_to_string(page.to_image(resolution=180).original) for page in pdf.pages)
    else: text = pytesseract.image_to_string(Image.open(path))
    return [pd.DataFrame([{"document_text": text}])]

def process_file(db: Session, file_id: int):
    f = db.get(UploadedFile, file_id)
    if not f: return
    f.status, f.progress = "processing", 1; db.commit()
    try:
        # Index prior project records once, then compare only records in matching blocks.
        from collections import defaultdict
        blocks=defaultdict(set)
        prior_rows=db.query(Record).filter(Record.project_id==f.project_id,Record.status.in_(["active","review"])).all()
        for prior in prior_rows:
            for key in block_keys(prior.data): blocks[key].add(prior.id)
        chunks = read_rows(Path(f.storage_path)); count = 0; mapping = None
        for chunk in chunks:
            chunk = chunk.where(pd.notna(chunk), None)
            if mapping is None:
                mapping=f.mapping or suggest_mapping(chunk.columns)
                if not f.mapping:
                    from app.services.llm import map_columns
                    mapping=map_columns([str(c) for c in chunk.columns],mapping)
            for item in chunk.to_dict(orient="records"):
                raw = {str(k): (None if pd.isna(v) else str(v)) for k, v in item.items()}
                source_raw = raw
                if "document_text" in raw:
                    from app.services.llm import extract_invoice
                    extracted, note = extract_invoice(raw.get("document_text") or "")
                    if extracted:
                        source_raw = extracted
                        mapping = {key:key for key in extracted}
                    else:
                        raw["_review_note"] = note
                data, changes, flags = clean_record(source_raw, mapping)
                if raw.get("_review_note"): flags.append(raw["_review_note"])
                rec = Record(project_id=f.project_id, file_id=f.id, row_number=count + 1, raw_data=raw, data=data, quality_flags=flags)
                db.add(rec); db.flush()
                for field, old, new, reason in changes:
                    db.add(AuditLog(project_id=f.project_id, record_id=rec.id, field=field, old_value=None if old is None else str(old), new_value=str(new), reason=reason))
                # Blocking by exact identifier or initial avoids all-pairs comparisons.
                candidate_ids=set().union(*(blocks[key] for key in block_keys(data))) if block_keys(data) else set()
                prior = [db.get(Record, candidate_id) for candidate_id in candidate_ids]
                best = (None, 0.0)
                for old_rec in prior:
                    if not old_rec or old_rec.status not in ("active","review"): continue
                    if should_compare(data, old_rec.data):
                        score = score_records(data, old_rec.data)
                        if score > best[1]: best = (old_rec, score)
                if best[0] and best[1] > 90:
                    target = best[0]; previous=dict(target.data); target.data = {k: target.data.get(k) or data.get(k) for k in set(target.data)|set(data)}
                    rec.status, rec.duplicate_of, rec.duplicate_score = "merged", target.id, best[1]
                    db.add(AuditLog(project_id=f.project_id, record_id=target.id, field=f"merge:{rec.id}", old_value=json.dumps(previous), new_value=json.dumps(target.data), reason=f"Auto-merged duplicate with {best[1]}% match"))
                elif best[0] and best[1] >= 70: rec.status, rec.duplicate_of, rec.duplicate_score = "review", best[0].id, best[1]
                if rec.status in ("active","review"):
                    for key in block_keys(data): blocks[key].add(rec.id)
                count += 1
            f.row_count = count; f.progress = min(95, 10 + count // 1000); db.commit()
        f.status, f.progress = "complete", 100; f.row_count = count; db.commit()
    except Exception as exc:
        db.rollback(); f = db.get(UploadedFile, file_id)
        if f: f.status, f.error, f.progress = "failed", "Could not process this file. Check that it is a valid supported format and try again.", 0; db.commit()
